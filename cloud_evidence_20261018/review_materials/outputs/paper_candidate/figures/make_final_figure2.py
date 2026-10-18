import base64, pathlib
src=pathlib.Path(r'D:/Local_Projects/airslot/outputs/paper_candidate/figures/figure2_controlled_information_decision_value_final_base.png')
data=base64.b64encode(src.read_bytes()).decode()
svg='<svg xmlns="http://www.w3.org/2000/svg" width="1936" height="801" viewBox="0 0 1936 801"><image width="1936" height="801" preserveAspectRatio="none" href="data:image/png;base64,'+data+'"/></svg>'
pathlib.Path(r'D:/Local_Projects/airslot/outputs/paper_candidate/figures/figure2_controlled_information_decision_value_final.svg').write_text(svg,encoding='utf-8')
from pptx import Presentation
from pptx.util import Inches
p=Presentation()
p.slide_width=Inches(13.333333)
p.slide_height=Inches(13.333333*801/1936)
s=p.slides.add_slide(p.slide_layouts[6])
s.shapes.add_picture(str(src),0,0,width=p.slide_width,height=p.slide_height)
p.save(r'D:/Local_Projects/airslot/outputs/paper_candidate/figures/figure2_controlled_information_decision_value_final.pptx')
